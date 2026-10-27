"""Automated operational reporting run: load -> data quality checks -> quarantine -> cross-source
reconciliation -> daily and weekly KPIs -> anomaly detection -> Tableau extracts, charts, dashboard, Excel."""
import sqlite3
import time
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from plotly.subplots import make_subplots

ROOT = Path(__file__).resolve().parents[1]
DATA, SQL, OUT = ROOT / "data", ROOT / "sql", ROOT / "outputs"
TABLEAU, CHARTS = OUT / "tableau", OUT / "charts"
ANOMALY_KPIS = ["avg_turnaround_days", "processing_rate_pct", "auto_adjudication_pct", "abandonment_pct", "asa_sec"]
CALL_KPIS = {"abandonment_pct", "asa_sec"}
Z_THRESHOLD, BASELINE_DAYS = 4.0, 28


def load(con):
    for table in ("claims_ops", "call_center", "daily_ops_summary"):
        pd.read_csv(DATA / f"{table}.csv").to_sql(table, con, index=False)


def q(con, name):
    return pd.read_sql_query((SQL / name).read_text(), con)


def quarantine(con):
    """Deduplicate, then move rows failing a row-level rule to quarantine tables instead of deleting them."""
    con.executescript("""
        CREATE TABLE claims_dedup AS SELECT DISTINCT * FROM claims_ops;
        CREATE TABLE claims_clean AS SELECT * FROM claims_dedup
         WHERE avg_turnaround_days IS NOT NULL AND avg_turnaround_days BETWEEN 0 AND 60
           AND received >= 0 AND processed >= 0 AND pended >= 0 AND denied >= 0 AND auto_adjudicated >= 0
           AND processed <= received * 1.3 AND denied <= processed AND auto_adjudicated <= processed;
        CREATE TABLE claims_quarantine AS SELECT * FROM claims_dedup EXCEPT SELECT * FROM claims_clean;
        CREATE TABLE calls_clean AS SELECT DISTINCT * FROM call_center
         WHERE avg_handle_time_sec IS NOT NULL AND calls_offered > 0
           AND calls_answered + abandoned = calls_offered;
        CREATE TABLE calls_quarantine AS SELECT * FROM call_center EXCEPT SELECT * FROM calls_clean;
    """)
    counts = {t: con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in (
        "claims_ops", "claims_dedup", "claims_clean", "claims_quarantine", "call_center", "calls_clean",
        "calls_quarantine")}
    return counts, pd.read_sql_query("SELECT * FROM claims_quarantine", con), \
        pd.read_sql_query("SELECT * FROM calls_quarantine", con)


def detect_anomalies(daily):
    """Rolling z-score of each KPI against the prior 28 days (current day excluded from its own baseline)."""
    d = daily.set_index(pd.to_datetime(daily.date))
    found = []
    for kpi in ANOMALY_KPIS:
        # incomplete days (a quarantined source row) would look like false anomalies, so they are skipped
        complete = d.queues_reported == 3 if kpi in CALL_KPIS else d.claim_types_reported == 3
        s = d.loc[complete, kpi]
        base = s.shift(1).rolling(BASELINE_DAYS, min_periods=14)
        z = (s - base.mean()) / base.std()
        hits = z[z.abs() > Z_THRESHOLD]
        found += [(i.date(), kpi, round(s[i], 2), round(base.mean()[i], 2), round(v, 2)) for i, v in hits.items()]
    an = pd.DataFrame(found, columns=["date", "kpi", "value", "baseline_28d_mean", "z_score"])
    return an.sort_values(["date", "kpi"]).reset_index(drop=True)


def group_events(an, gap_days=3):
    """Collapse anomaly days that are within a few days of each other into one event for the commentary."""
    if an.empty:
        return pd.DataFrame(columns=["event_start", "event_end", "days_flagged", "kpis", "max_abs_z"])
    d = an.assign(date=pd.to_datetime(an.date)).sort_values("date")
    d["event"] = (d.date.diff().dt.days.fillna(0) > gap_days).cumsum()
    return d.groupby("event").agg(
        event_start=("date", "min"), event_end=("date", "max"), days_flagged=("date", "nunique"),
        kpis=("kpi", lambda k: ", ".join(sorted(set(k)))), max_abs_z=("z_score", lambda z: round(z.abs().max(), 1)),
    ).reset_index(drop=True)


def explain_exceptions(ex, claims_q):
    """Label each reconciliation exception with its most likely cause for the research log."""
    bad_days = set(claims_q.date)

    def cause(r):
        if r.status == "MISSING IN SOURCE":
            return "Source extract missing for this date"
        if r.date in bad_days:
            return "Source row failed data quality checks (see quarantined_records.csv)"
        if r.received_var_pct < -0.5 and r.processed_var_pct < -0.5:
            return "Late-arriving source file not loaded to reporting system"
        if r.denied_var_pct and abs(r.denied_var_pct) > 0.5:
            return "Manual adjustment in reporting system not posted to source"
        return "Research with source system owner"

    ex["likely_cause"] = ex.apply(cause, axis=1)
    return ex


def dq_html(dq, counts, path):
    rows = "".join(
        f"<tr class='{'pass' if r.failed_rows == 0 else 'fail'}'><td>{r.source}</td><td>{r.dimension}</td>"
        f"<td>{r.check_name}</td><td>{r.failed_rows}</td><td>{'PASS' if r.failed_rows == 0 else 'FAIL'}</td></tr>"
        for r in dq.itertuples())
    summary = "".join(f"<li>{k}: {v:,} rows</li>" for k, v in counts.items())
    path.write_text(f"""<!doctype html><html><head><meta charset="utf-8"><title>Data Quality Report</title>
<style>body{{font-family:Segoe UI,Arial,sans-serif;margin:32px;color:#222}}h1{{color:#1F3864}}
table{{border-collapse:collapse;min-width:720px}}th,td{{border:1px solid #ccc;padding:6px 10px;text-align:left}}
th{{background:#1F3864;color:#fff}}tr.pass td:last-child{{background:#C6EFCE}}tr.fail td:last-child{{background:#F8CBAD}}
</style></head><body><h1>Operations Data Quality Report</h1>
<p>{(dq.failed_rows == 0).sum()} of {len(dq)} checks passed. Failing rows are quarantined, not deleted.</p>
<table><tr><th>Source</th><th>Dimension</th><th>Check</th><th>Failed rows</th><th>Status</th></tr>{rows}</table>
<h2>Record counts</h2><ul>{summary}</ul></body></html>""", encoding="utf-8")


def charts(daily, an):
    CHARTS.mkdir(parents=True, exist_ok=True)
    d = daily.assign(date=pd.to_datetime(daily.date))
    claim_cols = ["received", "processed", "avg_turnaround_days", "processing_rate_pct", "auto_adjudication_pct",
                  "denial_rate_pct"]
    call_cols = ["abandonment_pct", "answer_rate_pct", "asa_sec", "aht_sec"]
    d.loc[d.claim_types_reported < 3, claim_cols] = None  # incomplete days are not plotted
    d.loc[d.queues_reported < 3, call_cols] = None
    fig, ax = plt.subplots(1, 2, figsize=(13, 4.3))
    ax[0].plot(d.date, d.avg_turnaround_days, color="#1F3864", lw=0.8, label="Daily")
    ax[0].plot(d.date, d.avg_turnaround_days.rolling(7).mean(), color="#C55A11", lw=1.6, label="7-day avg")
    a = an[an.kpi == "avg_turnaround_days"]
    ax[0].scatter(pd.to_datetime(a.date), a.value, color="#C00000", zorder=3, label="Anomaly (|z| > 4)")
    ax[0].set(title="Claims turnaround (days)", ylabel="Avg turnaround days")
    ax[0].legend(fontsize=8)
    ax[1].plot(d.date, d.abandonment_pct, color="#1F3864", lw=0.8, label="Daily")
    ax[1].plot(d.date, d.abandonment_pct.rolling(7).mean(), color="#C55A11", lw=1.6, label="7-day avg")
    a = an[an.kpi == "abandonment_pct"]
    ax[1].scatter(pd.to_datetime(a.date), a.value, color="#C00000", zorder=3, label="Anomaly (|z| > 4)")
    ax[1].set(title="Call abandonment rate (%)", ylabel="% of calls offered")
    ax[1].legend(fontsize=8)
    for x in ax:
        x.xaxis.set_major_formatter(mdates.DateFormatter("%b %y"))
    fig.tight_layout()
    fig.savefig(CHARTS / "operations_kpis.png", dpi=130)
    plt.close(fig)

    fig = make_subplots(rows=2, cols=2, subplot_titles=(
        "Claims received vs processed", "Turnaround days (7-day avg)", "Abandonment % and answer rate %",
        "Auto-adjudication % and denial %"))
    fig.add_scatter(x=d.date, y=d.received, name="Received", row=1, col=1)
    fig.add_scatter(x=d.date, y=d.processed, name="Processed", row=1, col=1)
    fig.add_scatter(x=d.date, y=d.avg_turnaround_days.rolling(7).mean(), name="Turnaround 7d", row=1, col=2)
    fig.add_scatter(x=d.date, y=d.abandonment_pct, name="Abandonment %", row=2, col=1)
    fig.add_scatter(x=d.date, y=d.answer_rate_pct, name="Answer rate %", row=2, col=1)
    fig.add_scatter(x=d.date, y=d.auto_adjudication_pct, name="Auto-adjudication %", row=2, col=2)
    fig.add_scatter(x=d.date, y=d.denial_rate_pct, name="Denial %", row=2, col=2)
    for kpi, (r, c) in (("avg_turnaround_days", (1, 2)), ("abandonment_pct", (2, 1))):
        a = an[an.kpi == kpi]
        fig.add_scatter(x=pd.to_datetime(a.date), y=a.value, mode="markers", marker=dict(color="red", size=8),
                        name=f"Anomaly: {kpi}", row=r, col=c)
    fig.update_layout(height=820, template="plotly_white",
                      title="Operational Performance Dashboard (synthetic data)")
    fig.write_html(OUT / "operations_dashboard.html", include_plotlyjs="cdn")


def _sheet(wb, title, df):
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])
    t = Table(displayName=f"tbl{title.replace(' ', '')}", ref=f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(str(col)) + 2)
    ws.freeze_panes = "A2"
    return ws


def excel(weekly, dq, recon_ex, an, path):
    wb = Workbook()
    wb.remove(wb.active)
    ws = _sheet(wb, "Weekly KPIs", weekly)
    n = len(weekly) + 1
    cols = {c: get_column_letter(i + 1) for i, c in enumerate(weekly.columns)}
    ws.conditional_formatting.add(f"{cols['abandonment_pct']}2:{cols['abandonment_pct']}{n}",
                                  ColorScaleRule(start_type="min", start_color="C6EFCE",
                                                 end_type="max", end_color="F8CBAD"))
    ws.conditional_formatting.add(f"{cols['avg_turnaround_days']}2:{cols['avg_turnaround_days']}{n}",
                                  ColorScaleRule(start_type="min", start_color="C6EFCE",
                                                 end_type="max", end_color="F8CBAD"))
    ws = _sheet(wb, "Data Quality", dq)
    ws.conditional_formatting.add(f"D2:D{len(dq) + 1}", CellIsRule(operator="greaterThan", formula=["0"],
                                                                   fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(f"D2:D{len(dq) + 1}", CellIsRule(operator="equal", formula=["0"],
                                                                   fill=PatternFill("solid", fgColor="C6EFCE")))
    _sheet(wb, "Reconciliation", recon_ex)
    _sheet(wb, "Anomalies", an)
    wb.active = 0
    wb["Weekly KPIs"]["A1"].font = Font(bold=True, color="FFFFFF")
    wb.save(path)


def main():
    start = time.perf_counter()
    for p in (OUT, TABLEAU):
        p.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(":memory:")
    load(con)
    dq = q(con, "01_data_quality_checks.sql")
    counts, claims_q, calls_q = quarantine(con)
    recon = q(con, "02_cross_source_reconciliation.sql")
    daily = q(con, "03_daily_kpis.sql")
    daily.to_sql("daily_kpis", con, index=False)
    weekly = q(con, "04_weekly_kpi_trends.sql")
    an = detect_anomalies(daily)
    recon_ex = explain_exceptions(recon[recon.status != "OK"].copy(), claims_q)

    dq.to_csv(OUT / "data_quality_report.csv", index=False)
    dq_html(dq, counts, OUT / "data_quality_report.html")
    recon_ex.to_csv(OUT / "reconciliation_exceptions.csv", index=False)
    pd.concat([claims_q.assign(source="claims_ops"), calls_q.assign(source="call_center")]).to_csv(
        OUT / "quarantined_records.csv", index=False)
    an.to_csv(OUT / "anomalies.csv", index=False)
    events = group_events(an)
    events.to_csv(OUT / "anomaly_events.csv", index=False, date_format="%Y-%m-%d")
    daily.to_csv(TABLEAU / "daily_kpis.csv", index=False)
    weekly.to_csv(TABLEAU / "weekly_kpis.csv", index=False)
    pd.read_sql_query("SELECT * FROM claims_clean", con).to_csv(TABLEAU / "claims_by_type.csv", index=False)
    pd.read_sql_query("SELECT * FROM calls_clean", con).to_csv(TABLEAU / "calls_by_queue.csv", index=False)
    charts(daily, an)
    excel(weekly, dq, recon_ex, an, OUT / "Operations_KPI_Report.xlsx")
    elapsed = time.perf_counter() - start

    print(dq.to_string(index=False), "\n")
    print("Record counts:", counts, "\n")
    print(recon_ex[["date", "received_var_pct", "processed_var_pct", "denied_var_pct", "likely_cause"]].to_string(index=False),
          "\n")
    print(an.to_string(index=False), "\n")
    print(events.to_string(index=False), "\n")
    rows = counts["claims_ops"] + counts["call_center"] + len(recon)
    print(f"Manual review replaced: {len(dq)} data quality checks + {len(recon)}-day reconciliation + "
          f"{len(ANOMALY_KPIS)} KPI anomaly scans across {rows:,} rows in {elapsed:.1f} s")


if __name__ == "__main__":
    main()
