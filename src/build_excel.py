"""Build the Excel membership report: data tables, a self-service report with dropdowns, SUMIFS / XLOOKUP /
INDEX-MATCH / VLOOKUP formulas, budget variance with conditional formatting, and a PivotTable.
The PivotTable is added through Excel itself (Windows + Excel required); without Excel the rest still builds."""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F3864")
RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="C6EFCE")


def _table(ws, df, name, number_formats=None):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])
    ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(col) + 2)
        if number_formats and col in number_formats:
            for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell[0].number_format = number_formats[col]
    ws.freeze_panes = "A2"


def _report_sheet(wb, products, months):
    ws = wb.create_sheet("Report", 0)
    ws["A1"] = "Membership Report - select a product and month"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    labels = ["Product", "Month", "Members", "Members, same month last year", "YoY change", "YoY %",
              "Gross adds", "Terminations", "Line of business (INDEX-MATCH)", "Owner team (XLOOKUP)",
              "Owner team (VLOOKUP)", "2026 budget", "Variance to budget %"]
    for i, lab in enumerate(labels, start=3):
        ws.cell(row=i, column=1, value=lab).font = Font(bold=True)
    ws["B3"], ws["B4"] = products[0], months[-1]
    ws["B4"].number_format = "mmm yyyy"
    f = {
        5: "=SUMIFS(Data!$D:$D,Data!$B:$B,$B$3,Data!$A:$A,$B$4)",
        6: "=SUMIFS(Data!$D:$D,Data!$B:$B,$B$3,Data!$A:$A,EDATE($B$4,-12))",
        7: "=B5-B6",
        8: '=IFERROR(B7/B6,"n/a")',
        9: "=SUMIFS(Data!$E:$E,Data!$B:$B,$B$3,Data!$A:$A,$B$4)",
        10: "=SUMIFS(Data!$F:$F,Data!$B:$B,$B$3,Data!$A:$A,$B$4)",
        11: "=INDEX(Products!$B:$B,MATCH($B$3,Products!$A:$A,0))",
        12: '=_xlfn.XLOOKUP($B$3,Products!$A:$A,Products!$C:$C,"Not found")',
        13: "=VLOOKUP($B$3,Products!$A:$C,3,FALSE)",
        14: "=SUMIFS(Budget!$D:$D,Budget!$B:$B,$B$3,Budget!$A:$A,$B$4)",
        15: '=IF(B14=0,"no budget",(B5-B14)/B14)',
    }
    for r, formula in f.items():
        ws.cell(row=r, column=2, value=formula)
    for r in (5, 6, 7, 9, 10, 14):
        ws.cell(row=r, column=2).number_format = "#,##0"
    for r in (8, 15):
        ws.cell(row=r, column=2).number_format = "0.0%"
    dv_p = DataValidation(type="list", formula1=f"=Products!$A$2:$A${len(products) + 1}", allow_blank=False)
    dv_m = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${len(months) + 1}", allow_blank=False)
    dv_p.error, dv_m.error = "Pick a product from the list", "Pick a month from the list"
    ws.add_data_validation(dv_p)
    ws.add_data_validation(dv_m)
    dv_p.add("B3")
    dv_m.add("B4")
    for c in ("B3", "B4"):
        ws[c].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.conditional_formatting.add("B8", CellIsRule(operator="lessThan", formula=["0"], fill=RED))
    ws.conditional_formatting.add("B15", FormulaRule(formula=["AND(ISNUMBER(B15),ABS(B15)>0.02)"], fill=RED))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 36, 18
    ws["D3"] = "Yellow cells are dropdowns (data validation). All other values are live formulas."
    ws["D3"].font = Font(italic=True, color="7F7F7F")
    ws["D12"] = "XLOOKUP needs Excel 2021 or Microsoft 365; VLOOKUP and INDEX-MATCH give the same answer."
    ws["D12"].font = Font(italic=True, color="7F7F7F")


def build_workbook(path, summary, budget_var, fc, recon, con):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    data = summary.copy()
    data["month"] = pd.to_datetime(data.month)
    data["year"] = data.month.dt.year
    data = data[["month", "product", "line_of_business", "members", "gross_adds", "terminations",
                 "net_change", "mom_pct", "yoy_pct", "year"]]
    _table(ws, data, "tblMembership", {"month": "yyyy-mm", "members": "#,##0"})

    products = pd.read_sql_query("SELECT * FROM dim_product", con)
    _table(wb.create_sheet("Products"), products, "tblProducts")

    bv = budget_var.copy()
    bv["month"] = pd.to_datetime(bv.month)
    wsb = wb.create_sheet("Budget")
    _table(wsb, bv[["month", "product", "actual_members", "budget_members"]], "tblBudget",
           {"month": "yyyy-mm", "actual_members": "#,##0", "budget_members": "#,##0"})
    wsb["E1"], wsb["F1"] = "variance", "variance_pct"
    for r in range(2, len(bv) + 2):
        wsb[f"E{r}"] = f"=C{r}-D{r}"
        wsb[f"F{r}"] = f"=IFERROR(E{r}/D{r},0)"
        wsb[f"E{r}"].number_format, wsb[f"F{r}"].number_format = "#,##0", "0.0%"
    for c in ("E1", "F1"):
        wsb[c].font, wsb[c].fill = HEAD, FILL
    rng = f"F2:F{len(bv) + 1}"
    wsb.conditional_formatting.add(rng, FormulaRule(formula=["ABS(F2)>0.02"], fill=RED))
    wsb.conditional_formatting.add(rng, FormulaRule(formula=["ABS(F2)<=0.02"], fill=GREEN))
    wsb.column_dimensions["E"].width = wsb.column_dimensions["F"].width = 14

    f = fc.copy()
    f["month"] = pd.to_datetime(f.month)
    _table(wb.create_sheet("Forecast"), f, "tblForecast", {"month": "yyyy-mm"})
    _table(wb.create_sheet("Reconciliation"), recon[recon.status != "OK"], "tblRecon")

    months = sorted(data.month.unique())
    wl = wb.create_sheet("Lists")
    wl["A1"] = "month"
    for i, m in enumerate(months, start=2):
        wl.cell(row=i, column=1, value=pd.Timestamp(m).to_pydatetime()).number_format = "mmm yyyy"
    wl.sheet_state = "hidden"

    _report_sheet(wb, list(products["product"]), [pd.Timestamp(m).to_pydatetime() for m in months])
    wb.save(path)
    _add_pivot(Path(path))


def _add_pivot(path):
    try:
        import win32com.client as win32
    except ImportError:
        print("pywin32 not available - skipped PivotTable")
        return
    xl = win32.DispatchEx("Excel.Application")
    xl.Visible, xl.DisplayAlerts = False, False
    try:
        wb = xl.Workbooks.Open(str(path.resolve()))
        src = wb.Worksheets("Data").ListObjects("tblMembership").Range
        ws = wb.Worksheets.Add()
        ws.Name = "Pivot"
        ws.Move(Before=wb.Worksheets("Data"))  # Report, Pivot, Data, ...
        cache = wb.PivotCaches().Create(SourceType=1, SourceData=src)
        pt = cache.CreatePivotTable(TableDestination=ws.Range("A3"), TableName="ptMembership")
        pt.PivotFields("line_of_business").Orientation = 3  # page filter
        pt.PivotFields("product").Orientation = 1  # rows
        pt.PivotFields("year").Orientation = 2  # columns
        df = pt.AddDataField(pt.PivotFields("members"), "Avg monthly members", -4106)  # xlAverage
        df.NumberFormat = "#,##0"
        ws.Range("A1").Value = "PivotTable: average monthly members by product and year (filter by line of business)"
        ws.Range("A1").Font.Bold = True
        wb.Worksheets("Report").Activate()
        xl.CalculateFull()
        wb.Save()
        wb.Close()
    finally:
        xl.Quit()
