"""Build a short PowerPoint deck that turns the analysis into insights for business stakeholders."""
from pptx import Presentation
from pptx.util import Inches, Pt

NAVY = (0x1F, 0x38, 0x64)


def _title_slide(prs, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.placeholders[1].text = subtitle


def _bullets(prs, title, items):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.text = items[0]
    for it in items[1:]:
        tf.add_paragraph().text = it
    for p in tf.paragraphs:
        p.font.size = Pt(18)


def build_deck(path, trend, recap, prev, dq, excluded, charts_dir):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _title_slide(prs, "Risk Adjustment Accuracy Review",
                 "Medicare Advantage & ACA | Payment years 2023-2025 | Synthetic data demo")

    latest = trend.payment_year.max()
    lt = trend[trend.payment_year == latest]
    top = lt.sort_values("avg_raf", ascending=False).iloc[0]
    lr = recap[recap.payment_year == latest].sort_values("recapture_rate_pct")
    low, high = lr.iloc[0], lr.iloc[-1]
    ma = lt[lt.segment == "MA"]
    _bullets(prs, "Key insights", [
        f"Highest average RAF in {latest}: {top.plan} ({top.avg_raf:.3f}), {top.raf_yoy_pct:+.1f}% vs prior year.",
        f"MA average RAF ranges {ma.avg_raf.min():.3f}-{ma.avg_raf.max():.3f} across plans.",
        f"Lowest HCC recapture: {low.plan} at {low.recapture_rate_pct:.1f}%; highest: {high.plan} at "
        f"{high.recapture_rate_pct:.1f}%.",
        "Each un-recaptured chronic HCC understates risk scores and revenue; recapture gaps are the "
        "largest controllable driver of RAF accuracy.",
    ])

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "RAF trend and HCC recapture by plan"
    s.shapes.add_picture(str(charts_dir / "raf_trend_and_recapture.png"), Inches(0.6), Inches(1.6), width=Inches(12))

    top_hcc = prev[(prev.payment_year == latest) & (prev.segment == "MA")].head(5)
    _bullets(prs, f"Most prevalent MA conditions, {latest}",
             [f"{r.hcc} {r.description}: {r.prevalence_pct:.1f}% of members (coef {r.coefficient})"
              for r in top_hcc.itertuples()])

    _bullets(prs, "Data quality and method", [
        f"{len(dq)} validation checks run before reporting; {excluded} records quarantined "
        f"({', '.join(dq[dq.failed_rows > 0].check_name)}).",
        "Recapture rate = prior-year HCCs documented again this year / prior-year HCCs, per member and HCC.",
        "YoY RAF change tested with Welch's t-test; RAF drivers estimated with OLS regression.",
        "Data is synthetic and CMS-structured; HCC coefficients are illustrative, not official CMS rates.",
    ])
    _bullets(prs, "Recommended next steps", [
        f"Prioritize chart reviews and annual wellness visits for {low.plan} members with un-recaptured HCCs.",
        "Add a monthly recapture tracker to the operational dashboard by provider group.",
        "Validate RAF movement against finance and actuarial revenue forecasts each quarter.",
    ])
    prs.save(path)
